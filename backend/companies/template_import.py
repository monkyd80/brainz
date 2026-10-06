"""Stable, cell-addressed site registration workbook.

The source sample is a presentation form with merged cells and ambiguous labels.
This workbook keeps its business sections but gives every input a unique key.
"""
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from .excel_dates import parse_excel_date


SECTIONS = [
    ("사이트 정보", [("name", "사이트명 *"), ("sales_manager", "담당 영업"), ("support_staff", "지원 인력"), ("contract_type", "계약 구분"), ("build_date", "구축 날짜 (YYYY-MM-DD)"), ("business_code", "사업 코드"), ("maintenance_code", "유지보수 코드"), ("current_file_date", "현행화 일자 (YYYY-MM-DD)"), ("supplied_server_os_db", "Server/OS/DB 납품 여부 (예/아니요)"), ("remote_support", "원격 지원 여부 (예/아니요)"), ("address", "위치")]),
    ("계약·제품", [("ems_package", "EMS 패키지"), ("ems_version", "EMS 버전"), ("dashboard_package", "Dashboard 패키지"), ("dashboard_version", "Dashboard 버전"), ("datamanager_version", "Datamanager 버전"), ("dashboard_location", "Dashboard 설치 위치")]),
    ("담당자", [("contact_name", "담당자명"), ("contact_phone", "유선번호"), ("contact_mobile", "휴대전화"), ("contact_email", "E-mail"), ("contact_fax", "FAX"), ("inspection_report_submitted", "점검지 제출 여부 (예/아니요)"), ("inspection_report_detail", "점검지 제출 상세")]),
    ("솔루션 라이선스", [(f"license_{code}_{part}", f"{label} {caption}") for code, label in (("sms", "SMS"), ("nms", "NMS"), ("dbms", "DBMS"), ("apm", "APM"), ("syslog", "SYSLOG"), ("trap", "TRAP"), ("oz", "OZ 보고서")) for part, caption in (("used", "사용 개수"), ("total", "전체 개수"), ("other", "기타"), ("additional", "추가 정보"))]),
    ("연동 개발", [("integration_sms", "단문자 사용 (예/아니요)"), ("integration_sms_detail", "단문자 상세"), ("integration_email", "메일 사용 (예/아니요)"), ("integration_email_detail", "메일 상세"), ("integration_push", "PushAPP 사용 (예/아니요)"), ("integration_push_detail", "PushAPP 상세"), ("integration_other", "기타 연동 상세")]),
    ("시스템 구성", [("system_name", "시스템명"), ("system_set", "Set 구성"), ("system_os", "OS"), ("system_os_version", "OS 버전"), ("system_db", "DB"), ("system_db_version", "DB 버전"), ("system_ip", "IP Address"), ("system_access", "시스템 접속 정보"), ("web_access", "Web 접속 정보"), ("web_port", "Web 포트"), ("db_port", "DB 포트"), ("dashboard_port", "Dashboard 포트")]),
    ("설치 위치", [(f"install_{code}", label) for code, label in (("db", "DB"), ("datafile", "Zenius Datafile"), ("java", "JAVA"), ("tomcat", "TOMCAT"), ("manager", "Manager"), ("log", "LOG"))]),
    ("H/W 구성", [(f"{target}_{part}", f"{label} {caption}") for target, label in (("manager1", "Manager #1"), ("manager2", "Manager #2"), ("db", "DB")) for part, caption in (("cpu", "CPU"), ("memory", "Memory"), ("disk", "DISK"))] + [("hw_other", "H/W 기타")]),
    ("기타", [("notes", "참고 사항"), ("special_notes", "사이트 특이사항 및 전달사항")]),
]
FIELDS = {key: label for _, items in SECTIONS for key, label in items}
BOOLEAN_FIELDS = {"supplied_server_os_db", "remote_support", "inspection_report_submitted", "integration_sms", "integration_email", "integration_push"}
INTEGER_FIELDS = {"db_port"} | {f"license_{code}_{part}" for code in ("sms", "nms", "dbms", "apm", "syslog", "trap", "oz") for part in ("used", "total")}
DATE_FIELDS = {"build_date", "current_file_date"}


def template_bytes():
    book = Workbook()
    sheet = book.active
    sheet.title = "사이트 등록"
    sheet.append(["항목 코드 (수정 금지)", "입력 항목", "입력값"])
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(color="FFFFFF", bold=True)
    for section, items in SECTIONS:
        sheet.append([f"[{section}]", "", ""])
        for cell in sheet[sheet.max_row]:
            cell.fill = PatternFill("solid", fgColor="DCEAF7")
            cell.font = Font(bold=True)
        for key, label in items:
            sheet.append([key, label, None])
            sheet.cell(sheet.max_row, 3).fill = PatternFill("solid", fgColor="FFF2CC")
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 43
    sheet.column_dimensions["C"].width = 55
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = f"A1:C{sheet.max_row}"
    sheet["E1"] = "사용 방법"
    sheet["E2"] = "노란색 입력값 열만 작성하세요. 사이트명은 필수입니다."
    sheet["E3"] = "날짜: YYYY-MM-DD / 여부: 예 또는 아니요 / 수량·포트: 숫자"
    sheet["E4"] = "항목 코드와 시트 이름은 변경하지 마세요."
    sheet.column_dimensions["E"].width = 70
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def _convert(key, value):
    if value is None or str(value).strip() in ("", "-"):
        return None
    if key in BOOLEAN_FIELDS:
        normalized = str(value).strip().lower()
        if normalized in ("예", "유", "사용", "y", "yes", "true", "1"): return True
        if normalized in ("아니요", "무", "미사용", "n", "no", "false", "0"): return False
        raise ValueError(f"{FIELDS[key]}: 예 또는 아니요를 입력하세요.")
    if key in INTEGER_FIELDS:
        try:
            number = int(value)
            if number < 0: raise ValueError()
            return number
        except (ValueError, TypeError):
            raise ValueError(f"{FIELDS[key]}: 0 이상의 정수를 입력하세요.")
    if key in DATE_FIELDS:
        return parse_excel_date(value, FIELDS[key])
    return str(value).strip()


def preview_template(upload):
    if not upload.name.lower().endswith(".xlsx"):
        raise ValueError("새 등록 양식은 .xlsx 파일입니다.")
    try:
        sheet = load_workbook(BytesIO(upload.read()), data_only=True, read_only=True).active
    except Exception as error:
        raise ValueError("엑셀 파일을 읽을 수 없습니다.") from error
    if sheet.title != "사이트 등록" or sheet.cell(1, 1).value != "항목 코드 (수정 금지)":
        from .importers import workbook_grid, fixed_legacy_preview
        upload.seek(0)
        sheet_name, grid = workbook_grid(upload)
        legacy = fixed_legacy_preview(sheet_name, grid)
        if legacy:
            return legacy
        raise ValueError("사이트 등록 양식을 내려받아 입력한 파일을 선택하세요.")
    data = {}
    seen = set()
    for row in sheet.iter_rows(min_row=2, max_col=3, values_only=True):
        key = row[0]
        if key in FIELDS:
            if key in seen: raise ValueError(f"중복된 항목 코드: {key}")
            seen.add(key)
            data[key] = _convert(key, row[2])
    if seen != FIELDS.keys():
        raise ValueError("양식의 항목 코드가 변경되었거나 누락되었습니다. 새 양식을 내려받으세요.")
    mapped = [{"field": key, "label": FIELDS[key], "value": str(value)} for key, value in data.items() if value is not None]
    missing = [{"field": key, "label": FIELDS[key], "reason": "입력값 없음"} for key, value in data.items() if value is None]
    return {"sheet_name": sheet.title, "data": data, "mapped": mapped, "missing": missing, "unmapped_source_labels": []}
