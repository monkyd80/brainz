from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from .models import Company
from .template_import import template_bytes


class SiteTemplateImportTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="template_admin", password="test-password", is_staff=True)
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def upload(self, values):
        book = load_workbook(BytesIO(template_bytes()))
        sheet = book.active
        for row in sheet.iter_rows(min_row=2):
            if row[0].value in values:
                row[2].value = values[row[0].value]
        output = BytesIO()
        book.save(output)
        return SimpleUploadedFile("site_registration_template.xlsx", output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def test_download_preview_and_commit_store_related_details(self):
        download = self.client.get("/api/companies/import-template/")
        self.assertEqual(download.status_code, 200)
        self.assertEqual(load_workbook(BytesIO(download.content)).active.title, "사이트 등록")
        values = {"name": "테스트 사이트", "build_date": "2024-03-12", "ems_package": "EMS Plus", "contact_name": "홍길동", "license_sms_used": 0, "license_sms_total": 10, "integration_email": "예", "integration_email_detail": "알림", "system_name": "운영 서버", "install_db": "/data/db", "manager1_cpu": "8 core"}
        preview = self.client.post("/api/companies/import-preview/", {"file": self.upload(values)}, format="multipart")
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.data["data"]["name"], "테스트 사이트")
        self.assertEqual(preview.data["unmapped_source_labels"], [])
        created = self.client.post("/api/companies/import-commit/", {"file": self.upload(values)}, format="multipart")
        self.assertEqual(created.status_code, 201, created.data)
        company = Company.objects.get(name="테스트 사이트")
        self.assertEqual(company.build_year, 2024)
        self.assertEqual(company.contract_products.count(), 1)
        self.assertEqual(company.contacts.count(), 1)
        self.assertEqual(company.solutions.get(category="sms").used_quantity, 0)
        self.assertEqual(company.integrations.count(), 1)
        self.assertEqual(company.systems.count(), 1)
        self.assertEqual(company.installation_locations.count(), 1)
        self.assertEqual(company.hardware_specs.count(), 1)

    def test_invalid_boolean_does_not_create_company(self):
        response = self.client.post("/api/companies/import-commit/", {"file": self.upload({"name": "잘못된 사이트", "remote_support": "아마도"})}, format="multipart")
        self.assertEqual(response.status_code, 400)
        self.assertFalse(Company.objects.filter(name="잘못된 사이트").exists())
