import re
from io import BytesIO

import xlrd
from openpyxl import load_workbook
from .excel_dates import parse_excel_date


FIELD_MAP = {
    "name": ("사이트명", "사이트명"), "build_year": ("구축년도", "구축년도"),
    "business_code": ("사업/코드", "사업 코드"), "maintenance_code": ("유지보수코드", "유지보수 코드"),
    "contract_type": ("계약구분", "계약 구분"), "sales_manager": ("담당영업", "담당 영업"),
    "support_staff": ("지원/인력", "지원 인력"), "address": ("위치", "위치"),
    "remote_support": ("원격유무", "원격 지원"), "notes": ("참고사항", "참고 사항"),
    "special_notes": ("사이트특이사항및전달사항", "특이사항 및 전달사항"),
}
HEADER_HINTS = ("정보", "패키지", "버전", "코드", "구분", "담당", "지원", "인력", "유무", "위치", "솔루션", "라이선스", "설치", "연동", "구성", "포트", "제출", "접속", "기타", "사항", "현행화")

# 기존 현행화 양식은 병합 셀이 많아 항목명 주변을 탐색하면 값이 누락될 수 있다.
# 양식 크기가 동일할 때에는 실제 입력 셀의 고정 좌표(0-based)를 우선 사용한다.
LEGACY_CELL_MAP = {
    "name": (1, 2), "sales_manager": (1, 8), "support_staff": (1, 13), "support_staff_secondary": (2, 13),
    "build_year": (4, 2), "business_code": (4, 6), "maintenance_code": (4, 9),
    "current_file_date": (4, 14), "contract_type": (6, 0),
    "ems_package": (6, 2), "ems_version": (6, 5),
    "supplied_server_os_db": (6, 11), "remote_support": (6, 14),
    "dashboard_package": (8, 2), "dashboard_version": (8, 5),
    "datamanager_version": (8, 8), "dashboard_location": (8, 11),
    "contact_name": (10, 0), "contact_phone": (10, 3), "contact_mobile": (10, 5),
    "contact_email": (10, 8), "contact_fax": (10, 11),
    "inspection_report_detail": (10, 14), "address": (11, 2),
    "integration_sms": (16, 3), "integration_email": (16, 8),
    "integration_push": (16, 11), "integration_other": (16, 14),
    "system_set": (18, 2), "system_os": (18, 4), "system_os_version": (18, 6),
    "system_db_version": (18, 7), "system_ip": (18, 9), "system_access": (18, 11),
    "web_access": (18, 13), "web_port": (20, 4), "db_port": (20, 9),
    "dashboard_port": (20, 13),
    "install_db": (22, 3), "install_datafile": (23, 3), "install_java": (24, 3),
    "install_tomcat": (25, 3), "install_manager": (26, 3), "install_log": (27, 3),
    "manager1_cpu": (23, 10), "manager1_memory": (24, 10), "manager1_disk": (25, 10),
    "manager2_cpu": (23, 12), "manager2_memory": (24, 12), "manager2_disk": (25, 12),
    "db_cpu": (23, 14), "db_memory": (24, 14), "db_disk": (25, 14),
    "notes": (29, 0), "special_notes": (31, 0),
}
LEGACY_LICENSE_CELLS = {
    "sms": (13, 2), "nms": (13, 4), "dbms": (13, 6),
    "apm": (13, 8), "trap": (13, 9), "syslog": (13, 10),
}

SOLUTION_CATEGORY_BY_NAME = {
    "sms": "sms", "nms": "nms", "dbms": "dbms", "apm": "apm",
    "syslog": "syslog", "trap": "trap", "oz": "oz",
}


def normalize(value):
    return re.sub(r"\s+", "", str(value or "")).lower()


def empty_excel_value(value):
    return None if isinstance(value, str) and value.strip() == "-" else value


def is_header(value):
    return any(hint in normalize(value) for hint in HEADER_HINTS)


def workbook_grid(upload):
    data = upload.read()
    name = upload.name.lower()
    if name.endswith(".xls") and not name.endswith(".xlsx"):
        book = xlrd.open_workbook(file_contents=data)
        sheet = book.sheet_by_index(0)
        grid = [[sheet.cell_value(r, c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
        for row_start, row_end, col_start, col_end in sheet.merged_cells:
            value = sheet.cell_value(row_start, col_start)
            for row in range(row_start, row_end):
                for col in range(col_start, col_end):
                    grid[row][col] = value
        return sheet.name, [[empty_excel_value(value) for value in row] for row in grid]
    if name.endswith(".xlsx"):
        sheet = load_workbook(BytesIO(data), data_only=True).active
        grid = [[sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)] for r in range(1, sheet.max_row + 1)]
        for merged in sheet.merged_cells.ranges:
            value = sheet.cell(merged.min_row, merged.min_col).value
            for row in range(merged.min_row - 1, merged.max_row):
                for col in range(merged.min_col - 1, merged.max_col): grid[row][col] = value
        return sheet.title, [[empty_excel_value(value) for value in row] for row in grid]
    raise ValueError(".xls 또는 .xlsx 파일만 업로드할 수 있습니다.")


def value_near_label(grid, label):
    for row_index, row in enumerate(grid):
        for col_index, cell in enumerate(row):
            if normalize(cell) != label: continue
            for value in row[col_index + 1 : min(len(row), col_index + 8)]:
                if value not in (None, "") and len(normalize(value)) > 1 and not is_header(value) and normalize(value) not in [item[0] for item in FIELD_MAP.values()]: return value
            for lower_row in grid[row_index + 1 : min(len(grid), row_index + 3)]:
                if col_index < len(lower_row) and lower_row[col_index] not in (None, "") and len(normalize(lower_row[col_index])) > 1 and not is_header(lower_row[col_index]): return lower_row[col_index]
    return None


def cell_value(grid, row, col):
    return empty_excel_value(grid[row][col]) if row < len(grid) and col < len(grid[row]) else None


def parse_solution_quantity(value):
    """수량만 숫자로 변환하고 같은 셀의 나머지 내용은 기타로 보존한다."""
    if empty_excel_value(value) in (None, ""):
        return {}
    text = str(value).strip()
    if isinstance(value, float) and value.is_integer() and value >= 0:
        return {"used_quantity": int(value)}
    number = r"(?:[0-9]{1,3}(?:,[0-9]{3})+|[0-9]+)"
    match = re.fullmatch(
        rf"(?P<used>{number})(?:[ \t]*/[ \t]*(?P<total>{number}))?(?:\s+(?P<detail>.*))?",
        text, re.DOTALL,
    )
    if not match:
        return {"other_detail": text}
    result = {"used_quantity": int(match["used"].replace(",", ""))}
    if match["total"] is not None:
        result["total_quantity"] = int(match["total"].replace(",", ""))
    if match["detail"] and match["detail"].strip():
        result["other_detail"] = match["detail"].strip()
    return result


def system_rows_from_grid(grid):
    columns = {"set_configuration": 2, "os_type": 4, "os_version": 6, "database_version": 7,
               "ip_address": 9, "system_access_info": 11, "web_access_info": 13}
    start, end = 18, 20
    aliases = {"set구성": "set_configuration", "os": "os_type", "버전": "os_version", "os버전": "os_version",
               "db": "database_type", "db버전": "database_version", "ipaddress": "ip_address",
               "시스템접속정보": "system_access_info", "web접속정보": "web_access_info"}
    for index, row in enumerate(grid[:35]):
        labels = [normalize(value) for value in row]
        if "set구성" in labels and "ipaddress" in labels:
            columns = {}
            for col, label in enumerate(labels):
                field = aliases.get(label)
                if field and field not in columns:
                    columns[field] = col
            start, end = index + 1, min(len(grid), 35)
            break
    rows = []
    for index in range(start, min(end, len(grid))):
        labels = [normalize(value) for value in grid[index]]
        if any(any(marker in label for marker in ("포트", "port", "설치위치", "h/w", "hw구성", "참고사항")) for label in labels):
            break
        values = {field: str(cell_value(grid, index, col)).strip() for field, col in columns.items()
                  if cell_value(grid, index, col) not in (None, "")}
        if values:
            rows.append({"name": values.get("set_configuration") or f"시스템 {len(rows) + 1}", **values})
    return rows


def fixed_legacy_preview(sheet_name, grid):
    if len(grid) < 32 or max((len(row) for row in grid), default=0) < 15:
        return None
    # 현행화 양식의 실제 입력 영역은 1~35행이며, 이후 행은 참고용 설명이다.
    # 파일에 참고 행이 더 있어도 데이터 매핑에는 사용하지 않는다.
    grid = grid[:35]
    port_row = next((index for index, row in enumerate(grid) if index >= 18 and
                     any(normalize(value) in ("web포트", "webport", "db포트", "dbport", "dashboard포트", "dashboardport") for value in row)), 20)
    data, mapped, missing = {}, [], []
    for field, coordinate in LEGACY_CELL_MAP.items():
        if coordinate[0] >= 20:
            coordinate = (coordinate[0] + port_row - 20, coordinate[1])
        value = cell_value(grid, *coordinate)
        if value in (None, ""):
            missing.append({"field": field, "label": FIELD_MAP.get(field, ("", field))[1], "reason": "고정 입력 셀이 비어 있음"})
            continue
        if field == "build_year":
            digits = re.search(r"\d{4}", str(value))
            value = int(digits.group()) if digits else None
        elif field == "current_file_date":
            value = parse_excel_date(value, "현행화 일자 (O5)")
        elif field == "remote_support":
            value = normalize(value) in ("y", "yes", "true", "1", "사용", "유", "가능")
        elif field == "supplied_server_os_db":
            value = normalize(value) not in ("n", "no", "false", "0", "미납품", "불가", "미사용")
        elif field.startswith("integration_"):
            value = normalize(value) in ("y", "yes", "true", "1", "사용", "유", "가능")
        if value not in (None, ""):
            data[field] = value
            mapped.append({"field": field, "label": FIELD_MAP.get(field, ("", field))[1], "value": str(value)})

    if data.get("support_staff_secondary"):
        data["support_staff"] = " / ".join(part for part in (str(data.get("support_staff", "")).strip(), str(data["support_staff_secondary"]).strip()) if part)
    if data.get("inspection_report_detail"):
        data["inspection_report_submitted"] = True

    systems = system_rows_from_grid(grid)
    if systems:
        # 기존 양식의 포트는 시스템 표 아래 별도 공통 행이므로 첫 시스템에 보존한다.
        for source, target in (("web_port", "web_port"), ("db_port", "database_port"), ("dashboard_port", "dashboard_port")):
            if data.get(source) not in (None, ""):
                systems[0][target] = data[source]
        data["systems"] = systems
        for index, system in enumerate(systems):
            mapped.append({"field": "system", "label": f"시스템 정보 {index + 1}", "value": " / ".join(str(value) for value in system.values())})

    # 솔루션 열 순서는 파일마다 달라질 수 있다. 13행의 실제 제목과 바로 아래
    # 14행의 값을 한 쌍으로 읽는다. 목록에 없는 제목도 product_name으로 보존한다.
    solution_rows = []
    if len(grid) > 13:
        for column, header in enumerate(grid[12]):
            normalized_header = normalize(header)
            value = cell_value(grid, 13, column)
            # 값이 있는 열만 솔루션으로 간주한다. 표 제목 셀은 아래 값이 비어 있으므로 제외된다.
            if not normalized_header or value in (None, ""):
                continue
            code = next((known for known in SOLUTION_CATEGORY_BY_NAME if known in normalized_header), None)
            product_name = str(header).strip()
            solution = {"product_name": product_name, "category": SOLUTION_CATEGORY_BY_NAME.get(code, "other")}
            solution.update(parse_solution_quantity(value))
            solution_rows.append(solution)
            mapped.append({"field": "solution", "label": f"구축 솔루션 · {product_name}", "value": str(value)})

    # 제목을 인식할 수 없는 구형 양식만 기존 고정 좌표를 사용한다.
    if not solution_rows:
        for code, coordinate in LEGACY_LICENSE_CELLS.items():
            value = cell_value(grid, *coordinate)
            if value in (None, ""):
                continue
            product_name = code.upper() if code != "oz" else "OZ 보고서"
            solution = {"product_name": product_name, "category": code}
            solution.update(parse_solution_quantity(value))
            solution_rows.append(solution)

    data["solutions"] = solution_rows
    # 기존 미리보기·호환 데이터도 유지한다.
    for solution in solution_rows:
        code = solution["category"]
        if code == "other":
            continue
        value = "/".join(str(solution[key]) for key in ("used_quantity", "total_quantity") if solution.get(key) is not None) or solution.get("other_detail", "")
        if value in (None, ""):
            continue
        if solution.get("used_quantity") is not None:
            data[f"license_{code}_used"] = solution["used_quantity"]
        if solution.get("total_quantity") is not None:
            data[f"license_{code}_total"] = solution["total_quantity"]
        if solution.get("other_detail"):
            data[f"license_{code}_other"] = solution["other_detail"]

    return {"sheet_name": sheet_name, "layout": "legacy_fixed", "data": data, "mapped": mapped, "missing": missing, "unmapped_source_labels": []}


def preview_upload(upload):
    sheet_name, grid = workbook_grid(upload)
    fixed_preview = fixed_legacy_preview(sheet_name, grid)
    if fixed_preview:
        return fixed_preview
    data, mapped, missing = {}, [], []
    for field, (label, display) in FIELD_MAP.items():
        value = value_near_label(grid, label)
        if value in (None, ""):
            missing.append({"field": field, "label": display, "reason": "엑셀에서 값을 찾지 못해 빈값으로 등록"})
            continue
        if field == "build_year":
            digits = re.search(r"\d{4}", str(value)); value = int(digits.group()) if digits else None
        if field == "remote_support": value = normalize(value) in ("y", "yes", "true", "1", "사용", "유", "가능")
        if value not in (None, ""): data[field] = value; mapped.append({"field": field, "label": display, "value": str(value)})
    labels = {label for label, _ in FIELD_MAP.values()}
    source_labels = sorted({normalize(cell) for row in grid for cell in row if isinstance(cell, str) and len(normalize(cell)) > 1})
    unmapped = [label for label in source_labels if label not in labels and any(key in label for key in ("정보", "코드", "버전", "담당", "설치", "구성", "연동", "라이선스"))]
    return {"sheet_name": sheet_name, "layout": "label_search", "data": data, "mapped": mapped, "missing": missing, "unmapped_source_labels": unmapped[:30]}
