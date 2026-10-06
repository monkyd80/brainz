from datetime import date
from io import BytesIO
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook
from rest_framework.test import APIClient
from .models import Company, CompanyAssignment, Solution, System


class ExcelUpdateTests(TestCase):
    def setUp(self):
        self.admin = get_user_model().objects.create_user(username="update_admin", is_staff=True)
        self.owner = get_user_model().objects.create_user(username="update_owner")
        self.company = Company.objects.create(name="동일 사이트", current_file_date=date(2026, 6, 12), created_by=self.owner, notes="기존 메모")
        CompanyAssignment.objects.create(company=self.company, user=self.owner)
        System.objects.create(company=self.company, name="기존 시스템")
        Solution.objects.create(company=self.company, category="trap", product_name="기존 TRAP", used_quantity=1)
        self.client = APIClient()
        self.client.force_authenticate(self.admin)

    def file(self, incoming=20260701, invalid_port=False):
        book = Workbook()
        sheet = book.active
        sheet.title = "현행화"
        sheet.cell(2, 3, "동일 사이트")
        if incoming is not None:
            sheet.cell(5, 15, incoming)
        sheet.cell(13, 3, "TMS")
        sheet.cell(14, 3, "0/100")
        sheet.cell(19, 3, "새 시스템")
        sheet.cell(19, 10, "*.*.*.32")
        if invalid_port:
            sheet.cell(21, 10, "잘못된 숫자")
        sheet.cell(30, 1, "새 메모")
        sheet.cell(32, 15, "-")
        output = BytesIO()
        book.save(output)
        return SimpleUploadedFile("legacy.xlsx", output.getvalue())

    def preview(self, incoming=20260701):
        response = self.client.post("/api/companies/import-preview/", {"file": self.file(incoming)}, format="multipart")
        self.assertEqual(response.status_code, 200)
        return response.data["duplicate"]

    def commit(self, confirm=False, stamp=None, incoming=20260701, invalid_port=False):
        body = {"file": self.file(incoming, invalid_port)}
        if confirm:
            body.update(confirm_update="true", expected_updated_at=stamp or self.company.updated_at.isoformat())
        return self.client.post("/api/companies/import-commit/", body, format="multipart")

    def test_preview_and_cancel_do_not_modify_data(self):
        duplicate = self.preview()
        self.assertTrue(duplicate["can_update"])
        self.assertEqual(duplicate["incoming_date"], "2026-07-01")
        self.assertEqual(self.commit().status_code, 409)
        self.company.refresh_from_db()
        self.assertEqual(self.company.current_file_date, date(2026, 6, 12))
        self.assertEqual(self.company.systems.get().name, "기존 시스템")

    def test_confirm_updates_existing_site_and_preserves_assignment(self):
        duplicate = self.preview()
        response = self.commit(True, duplicate["existing_updated_at"])
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["id"], self.company.pk)
        self.assertEqual(Company.objects.filter(name="동일 사이트").count(), 1)
        self.company.refresh_from_db()
        self.assertEqual(self.company.current_file_date, date(2026, 7, 1))
        self.assertEqual(self.company.notes, "새 메모")
        self.assertEqual(self.company.created_by, self.owner)
        self.assertEqual(list(self.company.assignments.values_list("user_id", flat=True)), [self.owner.pk])
        self.assertEqual(self.company.systems.get().name, "새 시스템")
        self.assertEqual(self.company.solutions.get().product_name, "TMS")
        self.assertTrue(self.company.change_history.filter(table_name="excel_upload", action="update").exists())

    def test_equal_older_missing_dates_and_wrong_permissions_are_blocked(self):
        for incoming in (20260612, 20260611, None):
            with self.subTest(incoming=incoming):
                self.assertFalse(self.preview(incoming)["can_update"])
                self.assertEqual(self.commit(True, incoming=incoming).status_code, 409)
        outsider = get_user_model().objects.create_user(username="update_outsider")
        self.client.force_authenticate(outsider)
        self.assertFalse(self.preview()["can_update"])
        self.assertEqual(self.commit(True).status_code, 409)
        self.client.force_authenticate(self.owner)
        self.assertTrue(self.preview()["can_update"])

    def test_stale_confirmation_is_rejected(self):
        stamp = self.preview()["existing_updated_at"]
        self.company.notes = "다른 사용자의 수정"
        self.company.save()
        self.assertEqual(self.commit(True, stamp).status_code, 409)
        self.company.refresh_from_db()
        self.assertEqual(self.company.notes, "다른 사용자의 수정")

    def test_failed_related_import_rolls_back_update_and_deletions(self):
        response = self.commit(True, invalid_port=True)
        self.assertEqual(response.status_code, 400)
        self.company.refresh_from_db()
        self.assertEqual(self.company.current_file_date, date(2026, 6, 12))
        self.assertEqual(self.company.notes, "기존 메모")
        self.assertEqual(self.company.systems.get().name, "기존 시스템")
        self.assertEqual(self.company.solutions.get().product_name, "기존 TRAP")
