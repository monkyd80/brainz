from datetime import date
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase
from openpyxl import Workbook
from rest_framework.test import APIClient

from .importers import fixed_legacy_preview, parse_solution_platform_quantities, preview_upload
from .models import Company, Solution


class PlatformQuantityParsingTests(SimpleTestCase):
    def test_numbers_optional_units_and_zero(self):
        cases = [
            ("sms", "Linux/Windows : 49\nUNIX : 6", {"linux_windows_quantity": 49, "unix_quantity": 6}),
            ("sms", "Linux/Windows : 49대\r\nUNIX : 0 개", {"linux_windows_quantity": 49, "unix_quantity": 0}),
            ("sms", "linux / windows：1,234 EA / unix: 5대", {"linux_windows_quantity": 1234, "unix_quantity": 5}),
            ("nms", "SNMP : 65개\nICMP : 20대\nSDN : 0", {"snmp_quantity": 65, "icmp_quantity": 20, "sdn_quantity": 0}),
            ("nms", "snmp: 1,000 licenses; icmp: 2 EA; sdn: 3건", {"snmp_quantity": 1000, "icmp_quantity": 2, "sdn_quantity": 3}),
            ("dbms", "Oracle : 4 License(s)", {"oracle_quantity": 4}),
        ]
        for category, value, expected in cases:
            with self.subTest(value=value):
                self.assertEqual(parse_solution_platform_quantities(value, category), expected)

    def test_empty_labels_dash_and_unrelated_categories_are_blank(self):
        for value in (None, "", "-", "Linux/Windows :\nUNIX :", "Linux/Windows : -\nUNIX : -"):
            with self.subTest(value=value):
                self.assertEqual(parse_solution_platform_quantities(value, "sms"), {})
        self.assertEqual(parse_solution_platform_quantities("SNMP :\nICMP : 0대\nSDN :", "nms"), {"icmp_quantity": 0})
        self.assertEqual(parse_solution_platform_quantities("Oracle: 5", "apm"), {})

    def test_invalid_quantities_are_rejected_with_cell_and_label(self):
        for value in ("-1", "1.5", "1,5", "10/20", "숫자아님", "1대 2개"):
            with self.subTest(value=value), self.assertRaisesRegex(ValueError, "C15 UNIX"):
                parse_solution_platform_quantities(f"UNIX: {value}", "sms", "C15")

    def test_details_without_headers_use_fixed_coordinates(self):
        grid = [[None] * 15 for _ in range(32)]
        grid[14][2] = "Linux/Windows: 12대\nUNIX: 0"
        grid[14][4] = "SNMP: 3개\nICMP:\nSDN:"
        grid[14][6] = "Oracle: 4"
        values = fixed_legacy_preview("현행화", grid)["data"]
        self.assertEqual(len(values["solutions"]), 3)
        self.assertEqual(values["license_sms_linux_windows_quantity"], 12)
        self.assertEqual(values["license_sms_unix_quantity"], 0)
        self.assertEqual(values["license_nms_snmp_quantity"], 3)
        self.assertEqual(values["license_dbms_oracle_quantity"], 4)

    def test_existing_xls_template_with_blank_labels_is_compatible(self):
        path = Path(settings.BASE_DIR).parent / "sample_kdy.xls"
        preview = preview_upload(SimpleUploadedFile(path.name, path.read_bytes()))
        self.assertEqual(preview["layout"], "legacy_fixed")
        self.assertFalse(any(key.endswith("_quantity") for key in preview["data"]))


class PlatformQuantityUploadTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(get_user_model().objects.create_user(username="platform_import_admin", is_staff=True))

    def upload(self, *, detail_only=False, blank=False, invalid=False, reordered=False):
        book = Workbook()
        sheet = book.active
        sheet.title = "현행화"
        sheet["C2"] = "상세 수량 업로드"
        sheet["O5"] = 20261027
        sheet["O32"] = "-"
        sheet.merge_cells("A13:B15")
        sheet["A13"] = "구축 솔루션\n(License)"
        columns = [
            (3, "SMS", "49/55\n기존 SMS 설명", "Linux/Windows : 49대\nUNIX : 0개"),
            (5, "NMS", "65/65", "SNMP : 65개\nICMP : 0\nSDN : 3 EA"),
            (7, "DBMS", "4/5", "Oracle : 4 License(s)"),
        ]
        if reordered:
            columns[0], columns[1] = (3, *columns[1][1:]), (5, *columns[0][1:])
        for col, title, main, detail in columns:
            for row in (13, 14, 15):
                sheet.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
            sheet.cell(13, col, title)
            if not detail_only:
                sheet.cell(14, col, main)
            if blank:
                detail = {"SMS": "Linux/Windows :\nUNIX :", "NMS": "SNMP :\nICMP :\nSDN :", "DBMS": "Oracle :"}[title]
            if invalid and title == "SMS":
                detail = "Linux/Windows : -1대\nUNIX :"
            sheet.cell(15, col, detail)
        output = BytesIO()
        book.save(output)
        return SimpleUploadedFile("platform_license.xlsx", output.getvalue())

    def test_actual_xls_template_maps_numeric_and_unit_values(self):
        content = (Path(settings.BASE_DIR).parent / "sample_kdy.xls").read_bytes()
        # Keep BIFF string lengths unchanged to reuse the real merged .xls form
        # without introducing another workbook-writing dependency.
        for original, replacement in (
            ("Linux/Windows :\nUNIX :", "Linux/Windows:1\nUNIX:2"),
            ("SNMP :\nICMP :\nSDN :", "SNMP:3EA\nICMP:\nSDN:"),
            ("Oracle :", "Oracle:4"),
        ):
            self.assertEqual(len(original), len(replacement))
            self.assertEqual(content.count(original.encode("ascii")), 1)
            content = content.replace(original.encode("ascii"), replacement.encode("ascii"))
        response = self.client.post("/api/companies/import-preview/", {
            "file": SimpleUploadedFile("platform_license.xls", content),
        }, format="multipart")
        self.assertEqual(response.status_code, 200, response.data)
        solutions = {row["category"]: row for row in response.data["data"]["solutions"]}
        self.assertEqual(solutions["sms"]["linux_windows_quantity"], 1)
        self.assertEqual(solutions["sms"]["unix_quantity"], 2)
        self.assertEqual(solutions["nms"]["snmp_quantity"], 3)
        self.assertNotIn("icmp_quantity", solutions["nms"])
        self.assertEqual(solutions["dbms"]["oracle_quantity"], 4)

    def test_preview_and_commit_map_all_fields_without_merged_duplicates(self):
        preview = self.client.post("/api/companies/import-preview/", {"file": self.upload()}, format="multipart")
        self.assertEqual(preview.status_code, 200, preview.data)
        self.assertEqual(len(preview.data["data"]["solutions"]), 3)
        mappings = {item["field"]: item["value"] for item in preview.data["mapped"]}
        self.assertEqual(mappings["license_sms_unix_quantity"], "0")
        self.assertEqual(mappings["license_nms_snmp_quantity"], "65")
        self.assertEqual(mappings["license_dbms_oracle_quantity"], "4")
        self.assertFalse(Company.objects.exists())
        response = self.client.post("/api/companies/import-commit/", {"file": self.upload()}, format="multipart")
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(len(response.data["solutions"]), 3)
        company = Company.objects.get(name="상세 수량 업로드")
        expected = {
            "sms": {"linux_windows_quantity": 49, "unix_quantity": 0},
            "nms": {"snmp_quantity": 65, "icmp_quantity": 0, "sdn_quantity": 3},
            "dbms": {"oracle_quantity": 4},
        }
        for code, fields in expected.items():
            saved = company.solutions.get(category=code)
            serialized = next(row for row in response.data["solutions"] if row["category"] == code)
            for field, value in fields.items():
                self.assertEqual(getattr(saved, field), value)
                self.assertEqual(serialized[field], value)
        sms = company.solutions.get(category="sms")
        self.assertEqual((sms.used_quantity, sms.total_quantity, sms.other_detail), (49, 55, "기존 SMS 설명"))

    def test_details_are_imported_without_main_counts(self):
        response = self.client.post("/api/companies/import-commit/", {"file": self.upload(detail_only=True)}, format="multipart")
        self.assertEqual(response.status_code, 201, response.data)
        company = Company.objects.get()
        self.assertEqual(company.solutions.count(), 3)
        sms = company.solutions.get(category="sms")
        self.assertIsNone(sms.used_quantity)
        self.assertIsNone(sms.total_quantity)
        self.assertEqual(sms.unix_quantity, 0)

    def test_blank_detail_labels_keep_counts_and_nullable_fields(self):
        response = self.client.post("/api/companies/import-commit/", {"file": self.upload(blank=True)}, format="multipart")
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(Company.objects.get().solutions.count(), 3)
        self.assertIsNone(Company.objects.get().solutions.get(category="sms").unix_quantity)
        self.assertEqual(Company.objects.get().solutions.get(category="sms").used_quantity, 49)

    def test_blank_detail_only_does_not_create_empty_solutions(self):
        response = self.client.post("/api/companies/import-commit/", {"file": self.upload(detail_only=True, blank=True)}, format="multipart")
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data["solutions"], [])

    def test_reordered_headers_keep_details_with_the_correct_solution(self):
        response = self.client.post("/api/companies/import-commit/", {"file": self.upload(reordered=True)}, format="multipart")
        self.assertEqual(response.status_code, 201, response.data)
        company = Company.objects.get()
        self.assertEqual(company.solutions.count(), 3)
        self.assertEqual(company.solutions.get(category="sms").linux_windows_quantity, 49)
        self.assertEqual(company.solutions.get(category="nms").snmp_quantity, 65)

    def test_invalid_detail_is_a_clear_error_and_does_not_create_site(self):
        for endpoint in ("import-preview", "import-commit"):
            response = self.client.post(f"/api/companies/{endpoint}/", {"file": self.upload(invalid=True)}, format="multipart")
            self.assertEqual(response.status_code, 400, response.data)
            self.assertIn("C15 Linux/Windows", response.data["detail"])
        self.assertFalse(Company.objects.exists())

    def test_confirmed_update_replaces_details_on_existing_site(self):
        company = Company.objects.create(name="상세 수량 업로드", current_file_date=date(2026, 10, 26))
        Solution.objects.create(company=company, category="sms", linux_windows_quantity=999)
        preview = self.client.post("/api/companies/import-preview/", {"file": self.upload()}, format="multipart")
        self.assertEqual(preview.status_code, 200, preview.data)
        response = self.client.post("/api/companies/import-commit/", {
            "file": self.upload(), "confirm_update": "true",
            "expected_updated_at": preview.data["duplicate"]["existing_updated_at"],
        }, format="multipart")
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(company.solutions.count(), 3)
        self.assertEqual(company.solutions.get(category="sms").linux_windows_quantity, 49)
        self.assertEqual(company.solutions.get(category="sms").unix_quantity, 0)

    def test_invalid_update_keeps_existing_site_and_solution(self):
        company = Company.objects.create(name="상세 수량 업로드", current_file_date=date(2026, 10, 26))
        solution = Solution.objects.create(company=company, category="sms", linux_windows_quantity=999)
        response = self.client.post("/api/companies/import-commit/", {
            "file": self.upload(invalid=True), "confirm_update": "true",
            "expected_updated_at": company.updated_at.isoformat(),
        }, format="multipart")
        self.assertEqual(response.status_code, 400, response.data)
        company.refresh_from_db()
        solution.refresh_from_db()
        self.assertEqual(company.current_file_date, date(2026, 10, 26))
        self.assertEqual(solution.linux_windows_quantity, 999)
